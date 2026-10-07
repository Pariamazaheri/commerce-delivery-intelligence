"""Optional public Samand market snapshot; bounded, sequential acquisition."""

import argparse
import hashlib
import json
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse
from urllib.robotparser import RobotFileParser

import pandas as pd
import requests
from bs4 import BeautifulSoup

BASE = "https://bama.ir"
AGENT = "SamandMarketResearch/1.0"
SEEDS = [
    "/car/samand",
    "/car/samand-lx",
    "/car/samand-se",
    "/car/samand-soren",
    "/car/samand-x7",
    "/car/samand-el",
    "/car/samand-sarir",
    "/car/samand-lx-ef7",
    "/car/samand-lx-ef7cng",
    "/car/samand-lx-xu7",
    "/car/samand-soren-plusxu7p",
    "/car/samand-soren-pluscng",
    "/car/samand-soren-plustu5p",
    "/car/samand-soren-elxturbo",
]


def parse_listing(html: str) -> list[str]:
    """Extract unique public detail paths without reverse-engineered endpoints."""
    soup = BeautifulSoup(html, "html.parser")
    return list(
        dict.fromkeys(
            a["href"]
            for a in soup.select('a[href^="/car/detail-"]')
            if "-samand-" in a["href"]
        )
    )


def parse_detail(html: str, url: str) -> dict:
    """Resolve only the required fields from the observed Nuxt SSR reference pool."""
    soup = BeautifulSoup(html, "html.parser")
    script = soup.select_one("#__NUXT_DATA__")
    if script is None:
        raise ValueError("Missing server-rendered vehicle payload")
    pool = json.loads(script.string)

    def field(mapping, key):
        if not isinstance(mapping, dict):
            raise ValueError(f"Vehicle schema changed near {key}")
        reference = mapping.get(key, -1)
        if reference == -1:
            return None
        if not isinstance(reference, int) or not 0 <= reference < len(pool):
            raise ValueError(f"Invalid payload reference near {key}")
        return pool[reference]

    detail = next(
        (
            x
            for x in pool
            if isinstance(x, dict)
            and all(k in x for k in ["vehicle", "content", "price"])
        ),
        None,
    )
    if detail is None:
        raise ValueError("Vehicle schema changed")
    vehicle = field(detail, "vehicle")
    brand = field(vehicle, "brand")
    if field(brand, "value") != "samand":
        raise ValueError("Non-Samand vehicle")
    year = field(field(vehicle, "year"), "value")
    if not isinstance(year, int) or not 1386 <= year <= 1500:
        raise ValueError("Vehicle must be manufactured strictly after 1385")
    price = field(detail, "price")
    fixed = field(price, "fixed")
    price_value = field(fixed, "value") if fixed else None
    description = field(field(detail, "content"), "description") or ""
    # Do not republish phone numbers embedded in seller-written descriptions.
    description = re.sub(r"[0۰][9۹][\d۰-۹\s-]{9,}", "[phone removed]", description)
    row = {
        "price_toman": price_value,
        "price_type": field(price, "type"),
        "mileage_km": field(field(vehicle, "mileage"), "value"),
        "color": field(field(field(vehicle, "color"), "body"), "value"),
        "production_year_sh": year,
        "transmission": field(field(vehicle, "transmission"), "value"),
        "description": description,
        "source_url": url,
        "collected_at_utc": datetime.now(timezone.utc).isoformat(),
        "page_sha256": hashlib.sha256(html.encode()).hexdigest(),
    }
    if row["transmission"] not in ["دنده ای", "دنده‌ای", "اتوماتیک"]:
        raise ValueError("Unknown transmission type")
    row["transmission"] = "automatic" if row["transmission"] == "اتوماتیک" else "manual"
    if not isinstance(row["mileage_km"], (int, float)) or row["mileage_km"] < 0:
        raise ValueError("Invalid mileage")
    if not row["color"]:
        raise ValueError("Missing color")
    # Negotiable/installment prices remain missing rather than becoming zero.
    return row


def collect(output: Path, target: int = 50, pause: float = 1.5) -> dict:
    """Save real observations and an honest completion status, never padded rows."""
    output.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers["User-Agent"] = AGENT
    robots_response = session.get(f"{BASE}/robots.txt", timeout=30)
    robots_response.raise_for_status()
    robots_checked_at = datetime.now(timezone.utc).isoformat()
    robots = RobotFileParser()
    robots.parse(robots_response.text.splitlines())
    errors, rows, visited = [], [], set()

    def get(path):
        url = urljoin(BASE, path)
        if urlparse(url).hostname != "bama.ir" or not robots.can_fetch(AGENT, url):
            raise ValueError(f"Public crawl unavailable: {url}")
        time.sleep(pause)
        response = session.get(url, timeout=(15, 45))
        response.raise_for_status()
        return response.text

    for seed in SEEDS:
        if len(rows) >= target:
            break
        try:
            links = parse_listing(get(seed))
        except (requests.RequestException, ValueError) as error:
            errors.append({"url": urljoin(BASE, seed), "error": str(error)})
            continue
        for link in links:
            if len(rows) >= target:
                break
            if link in visited:
                continue
            visited.add(link)
            try:
                row = parse_detail(get(link), urljoin(BASE, link))
                rows.append(row)
                print(f"Collected {len(rows)}/{target}", flush=True)
            except (
                requests.RequestException,
                ValueError,
                TypeError,
                KeyError,
            ) as error:
                errors.append({"url": urljoin(BASE, link), "error": str(error)})
    pd.DataFrame(rows).to_csv(output / "samand_listings.csv", index=False)
    (output / "samand_listings.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    status = {
        "requested": target,
        "collected": len(rows),
        "complete": len(rows) == target,
        "sampling": "Latest visible public listings across Samand model filters",
        "errors": errors,
        "robots_checked_at_utc": robots_checked_at,
    }
    (output / "collection_status.json").write_text(json.dumps(status, indent=2))
    return status


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("reports/market_snapshot"))
    parser.add_argument("--target", type=int, default=50)
    parser.add_argument("--xlsx", action="store_true", help="Also export to Excel")
    args = parser.parse_args()
    if not 1 <= args.target <= 50:
        parser.error("target must be between 1 and 50")
    result = collect(args.output, args.target)
    if args.xlsx and result["collected"]:
        export_excel(args.output)
    print(json.dumps(result, indent=2))
    if not result["complete"]:
        raise SystemExit("Collection incomplete; see collection_status.json")


def export_excel(output: Path):
    """Portable runtime export; neutralize formula-like seller-authored text."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    rows = json.loads((output / "samand_listings.json").read_text(encoding="utf-8"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Samand listings"
    keys = list(rows[0])
    sheet.append(keys)
    for row in rows:
        sheet.append(
            [
                "'" + value
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
                else value
                for value in (row[key] for key in keys)
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="233044")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = 24
    sheet.column_dimensions["G"].width = 75
    sheet.column_dimensions["H"].width = 85
    for number in range(2, len(rows) + 2):
        sheet.cell(number, 1).number_format = "#,##0"
        sheet.cell(number, 3).number_format = "#,##0"
        description = rows[number - 2]["description"]
        lines = sum(max(1, (len(line) + 64) // 65) for line in description.split("\n"))
        sheet.row_dimensions[number].height = max(38, lines * 17 + 15)
        sheet.cell(number, 7).alignment = Alignment(wrap_text=True, vertical="top")
    table = Table(displayName="SamandListings", ref=f"A1:J{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    workbook.save(output / "samand_listings.xlsx")


if __name__ == "__main__":
    main()
